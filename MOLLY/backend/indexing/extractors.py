from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


# ============================================================
# Molly PTO — Document Extractors
# ============================================================


@dataclass
class ExtractedPart:
    """
    Часть документа:
    PDF page, Excel sheet, DOCX section/table и т.д.
    """

    part_type: str
    part_number: int | None = None
    name: str | None = None
    text: str = ""
    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass
class ExtractedDocument:
    """
    Унифицированный результат чтения любого файла.
    """

    text: str = ""

    parts: list[ExtractedPart] = field(
        default_factory=list
    )

    page_count: int = 0

    parser_name: str = ""

    parser_version: str = ""

    status: str = "ok"

    error: str | None = None

    is_scanned: bool = False

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


# ============================================================
# Общие функции
# ============================================================

def _safe_string(value: Any) -> str:
    if value is None:
        return ""

    return str(value)


def _clean_text(text: str) -> str:
    """
    Минимальная очистка без уничтожения структуры таблиц.
    """

    if not text:
        return ""

    text = text.replace(
        "\x00",
        " ",
    )

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{4,}",
        "\n\n\n",
        text,
    )

    return text.strip()


def _join_parts(
    parts: list[ExtractedPart],
) -> str:

    chunks: list[str] = []

    for part in parts:

        if part.name:
            chunks.append(
                f"[{part.part_type}: {part.name}]"
            )

        elif part.part_number is not None:
            chunks.append(
                f"[{part.part_type} {part.part_number}]"
            )

        if part.text:
            chunks.append(
                part.text
            )

    return "\n\n".join(
        chunks
    ).strip()


# ============================================================
# PDF
# ============================================================

def extract_pdf(
    path: Path,
) -> ExtractedDocument:

    try:
        from pypdf import PdfReader
    except ImportError as exc:

        return ExtractedDocument(
            status="error",
            error=f"pypdf не установлен: {exc}",
            parser_name="pypdf",
        )

    try:

        reader = PdfReader(
            str(path)
        )

        parts: list[ExtractedPart] = []

        total_text_length = 0

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):

            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = f"[Ошибка чтения страницы: {exc}]"

            text = _clean_text(text)

            total_text_length += len(text)

            parts.append(
                ExtractedPart(
                    part_type="page",
                    part_number=page_number,
                    name=f"Страница {page_number}",
                    text=text,
                )
            )

        full_text = _join_parts(parts)

        # PDF без текста при наличии страниц
        # считаем потенциально сканированным.
        is_scanned = (
            len(reader.pages) > 0
            and total_text_length < 100
        )

        metadata = {
            "pdf_page_count": len(reader.pages),
            "pdf_metadata": {
                str(key): str(value)
                for key, value in (
                    reader.metadata or {}
                ).items()
                if value is not None
            },
        }

        return ExtractedDocument(
            text=full_text,
            parts=parts,
            page_count=len(reader.pages),
            parser_name="pypdf",
            parser_version="6.x",
            status="ok",
            is_scanned=is_scanned,
            metadata=metadata,
        )

    except Exception as exc:

        return ExtractedDocument(
            status="error",
            error=str(exc),
            parser_name="pypdf",
        )


# ============================================================
# DOCX
# ============================================================

def extract_docx(
    path: Path,
) -> ExtractedDocument:

    try:
        from docx import Document
    except ImportError as exc:

        return ExtractedDocument(
            status="error",
            error=f"python-docx не установлен: {exc}",
            parser_name="python-docx",
        )

    try:

        document = Document(
            str(path)
        )

        parts: list[ExtractedPart] = []

        # ----------------------------------------------------
        # Основной текст
        # ----------------------------------------------------

        paragraph_lines: list[str] = []

        for paragraph in document.paragraphs:

            text = _clean_text(
                paragraph.text
            )

            if not text:
                continue

            style_name = ""

            try:
                if paragraph.style:
                    style_name = (
                        paragraph.style.name
                        or ""
                    )
            except Exception:
                pass

            if style_name:
                paragraph_lines.append(
                    f"[{style_name}] {text}"
                )
            else:
                paragraph_lines.append(
                    text
                )

        if paragraph_lines:

            parts.append(
                ExtractedPart(
                    part_type="text",
                    name="Основной текст",
                    text="\n".join(
                        paragraph_lines
                    ),
                )
            )

        # ----------------------------------------------------
        # Таблицы
        # ----------------------------------------------------

        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):

            rows: list[str] = []

            for row in table.rows:

                cells: list[str] = []

                for cell in row.cells:

                    cell_text = _clean_text(
                        cell.text
                    )

                    cells.append(
                        cell_text
                    )

                rows.append(
                    " | ".join(cells)
                )

            table_text = "\n".join(
                rows
            )

            parts.append(
                ExtractedPart(
                    part_type="table",
                    part_number=table_index,
                    name=f"Таблица {table_index}",
                    text=table_text,
                    metadata={
                        "rows": len(table.rows),
                        "columns": (
                            len(table.columns)
                            if table.rows
                            else 0
                        ),
                    },
                )
            )

        full_text = _join_parts(parts)

        # ----------------------------------------------------
        # Свойства документа
        # ----------------------------------------------------

        core = document.core_properties

        metadata = {
            "title": core.title or "",
            "subject": core.subject or "",
            "author": core.author or "",
            "keywords": core.keywords or "",
            "comments": core.comments or "",
            "category": core.category or "",
            "last_modified_by": (
                core.last_modified_by or ""
            ),
        }

        return ExtractedDocument(
            text=full_text,
            parts=parts,
            page_count=0,
            parser_name="python-docx",
            parser_version="1.x",
            status="ok",
            metadata=metadata,
        )

    except Exception as exc:

        return ExtractedDocument(
            status="error",
            error=str(exc),
            parser_name="python-docx",
        )


# ============================================================
# Excel
# ============================================================

def extract_excel(
    path: Path,
) -> ExtractedDocument:

    try:
        import openpyxl
    except ImportError as exc:

        return ExtractedDocument(
            status="error",
            error=f"openpyxl не установлен: {exc}",
            parser_name="openpyxl",
        )

    workbook = None

    try:

        # data_only=True: в индекс попадают значения ячеек,
        # а не формулы ("=Лист1!A5"). Если файл ни разу не
        # пересчитывался Excel'ом, значения формул будут пустыми.
        workbook = openpyxl.load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
            keep_links=False,
        )

        parts: list[ExtractedPart] = []

        total_cells = 0

        for sheet_number, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):

            lines: list[str] = []

            # Используем read_only режим.
            # Строки не загружаются целиком в память.
            # values_only=True: без создания объекта Cell на ячейку.
            for row in worksheet.iter_rows(
                values_only=True
            ):

                values: list[str] = []

                row_has_data = False

                for value in row:

                    if value is None:
                        values.append("")
                        continue

                    row_has_data = True
                    total_cells += 1

                    if isinstance(
                        value,
                        str,
                    ):
                        value_text = value
                    else:
                        value_text = str(
                            value
                        )

                    value_text = _clean_text(
                        value_text
                    )

                    values.append(
                        value_text
                    )

                if row_has_data:

                    lines.append(
                        " | ".join(values)
                    )

            sheet_text = "\n".join(
                lines
            )

            parts.append(
                ExtractedPart(
                    part_type="sheet",
                    part_number=sheet_number,
                    name=worksheet.title,
                    text=sheet_text,
                    metadata={
                        "max_row": worksheet.max_row,
                        "max_column": worksheet.max_column,
                    },
                )
            )

        sheet_count = (
            len(workbook.sheetnames)
            if hasattr(workbook, "sheetnames")
            else len(parts)
        )

        metadata = {
            "sheet_count": sheet_count,

            "total_cells": total_cells,

            "sheets": [
                part.name
                for part in parts
                if part.name
            ],
        }

        return ExtractedDocument(
            text=_join_parts(parts),
            parts=parts,
            page_count=0,
            parser_name="openpyxl",
            parser_version=getattr(
                openpyxl,
                "__version__",
                "",
            ),
            status="ok",
            metadata=metadata,
        )

    except Exception as exc:

        return ExtractedDocument(
            status="error",
            error=str(exc),
            parser_name="openpyxl",
        )

    finally:

        # На Windows незакрытая книга держит файл заблокированным
        # (в том числе при исключении посреди чтения).
        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


# ============================================================
# XLS
# ============================================================

def extract_xls(
    path: Path,
) -> ExtractedDocument:

    try:
        import xlrd
    except ImportError:

        return ExtractedDocument(
            status="unsupported",
            error=(
                "Для XLS требуется пакет xlrd."
            ),
            parser_name="xlrd",
        )

    try:

        workbook = xlrd.open_workbook(
            str(path),
            on_demand=True,
        )

        parts: list[ExtractedPart] = []

        total_cells = 0

        for sheet_number in range(
            workbook.nsheets
        ):

            sheet = workbook.sheet_by_index(
                sheet_number
            )

            lines: list[str] = []

            for row_index in range(
                sheet.nrows
            ):

                values: list[str] = []

                row_has_data = False

                for col_index in range(
                    sheet.ncols
                ):

                    value = sheet.cell_value(
                        row_index,
                        col_index,
                    )

                    if value in (
                        None,
                        "",
                    ):
                        values.append("")
                        continue

                    row_has_data = True
                    total_cells += 1

                    values.append(
                        _clean_text(
                            str(value)
                        )
                    )

                if row_has_data:
                    lines.append(
                        " | ".join(values)
                    )

            parts.append(
                ExtractedPart(
                    part_type="sheet",
                    part_number=sheet_number + 1,
                    name=sheet.name,
                    text="\n".join(lines),
                    metadata={
                        "rows": sheet.nrows,
                        "columns": sheet.ncols,
                    },
                )
            )

        workbook.release_resources()

        return ExtractedDocument(
            text=_join_parts(parts),
            parts=parts,
            parser_name="xlrd",
            parser_version=getattr(
                xlrd,
                "__version__",
                "",
            ),
            status="ok",
            metadata={
                "sheet_count": len(parts),
                "total_cells": total_cells,
                "sheets": [
                    part.name
                    for part in parts
                    if part.name
                ],
            },
        )

    except Exception as exc:

        return ExtractedDocument(
            status="error",
            error=str(exc),
            parser_name="xlrd",
        )


# ============================================================
# XML
# ============================================================

def _xml_local_name(tag: Any) -> str:

    if not isinstance(tag, str):
        return str(tag)

    if "}" in tag:
        return tag.split("}", 1)[1]

    return tag


def extract_xml(
    path: Path,
) -> ExtractedDocument:
    """
    Потоковое чтение XML (iterparse): дерево целиком в памяти
    не строится, обработанные элементы освобождаются, рекурсии нет
    (глубокая вложенность не даёт RecursionError).

    Формат вывода тот же, что был у ET.parse + walk:
    строки в порядке документа, отступ по уровню вложенности.
    """

    try:

        from config import MAX_DOCUMENT_TEXT

        lines: list[str] = []

        total_chars = 0

        root_tag = ""

        stack: list[ET.Element] = []

        # Элемент, чья строка ещё не выдана: его .text гарантированно
        # полон только к следующему событию парсера.
        pending: tuple[ET.Element, int] | None = None

        def emit(
            element: ET.Element,
            level: int,
        ) -> None:

            nonlocal total_chars

            text = (
                element.text or ""
            ).strip()

            attributes = " ".join(
                f"{key}={value}"
                for key, value in element.attrib.items()
            )

            if not (text or attributes):
                return

            line = "  " * level + _xml_local_name(
                element.tag
            )

            if attributes:
                line += f" [{attributes}]"

            if text:
                line += f": {text}"

            lines.append(line)

            total_chars += len(line) + 1

        for event, element in ET.iterparse(
            str(path),
            events=("start", "end"),
        ):

            if pending is not None:

                emit(*pending)

                pending = None

            if event == "start":

                if not stack:
                    root_tag = element.tag

                pending = (
                    element,
                    len(stack),
                )

                stack.append(element)

            else:

                stack.pop()

                # Освобождаем уже обработанный элемент.
                element.clear()

                if stack:
                    stack[-1].remove(element)

            # Дальше MAX_DOCUMENT_TEXT индексатор всё равно
            # обрежет текст — нет смысла читать гигабайты XML.
            if total_chars > MAX_DOCUMENT_TEXT:
                break

        if pending is not None:
            emit(*pending)

        text = _clean_text(
            "\n".join(lines)
        )

        return ExtractedDocument(
            text=text,
            parts=[
                ExtractedPart(
                    part_type="xml",
                    name=root_tag,
                    text=text,
                )
            ],
            parser_name="xml.etree",
            parser_version="stdlib",
            status="ok",
        )

    except Exception as exc:

        return ExtractedDocument(
            status="error",
            error=str(exc),
            parser_name="xml.etree",
        )


# ============================================================
# TXT
# ============================================================

def extract_txt(
    path: Path,
) -> ExtractedDocument:

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1251",
        "cp866",
    ]

    last_error: Exception | None = None

    for encoding in encodings:

        try:

            text = path.read_text(
                encoding=encoding,
                errors="strict",
            )

            text = _clean_text(
                text
            )

            return ExtractedDocument(
                text=text,
                parts=[
                    ExtractedPart(
                        part_type="text",
                        name="Текст",
                        text=text,
                    )
                ],
                parser_name=f"text:{encoding}",
                parser_version="stdlib",
                status="ok",
            )

        except Exception as exc:
            last_error = exc

    return ExtractedDocument(
        status="error",
        error=str(last_error),
        parser_name="text",
    )


# ============================================================
# CSV
# ============================================================

def extract_csv(
    path: Path,
) -> ExtractedDocument:

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1251",
    ]

    last_error: Exception | None = None

    for encoding in encodings:

        try:

            with path.open(
                "r",
                encoding=encoding,
                newline="",
            ) as file:

                sample = file.read(
                    8192
                )

                file.seek(0)

                try:
                    dialect = csv.Sniffer().sniff(
                        sample
                    )
                except csv.Error:
                    dialect = csv.excel

                reader = csv.reader(
                    file,
                    dialect,
                )

                lines: list[str] = []

                for row in reader:

                    lines.append(
                        " | ".join(
                            _clean_text(
                                str(value)
                            )
                            for value in row
                        )
                    )

            text = "\n".join(lines)

            return ExtractedDocument(
                text=text,
                parts=[
                    ExtractedPart(
                        part_type="table",
                        name="CSV",
                        text=text,
                    )
                ],
                parser_name=f"csv:{encoding}",
                parser_version="stdlib",
                status="ok",
            )

        except Exception as exc:
            last_error = exc

    return ExtractedDocument(
        status="error",
        error=str(last_error),
        parser_name="csv",
    )


# ============================================================
# JSON
# ============================================================

def extract_json(
    path: Path,
) -> ExtractedDocument:

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1251",
    ]

    last_error: Exception | None = None

    for encoding in encodings:

        try:

            with path.open(
                "r",
                encoding=encoding,
            ) as file:

                data = json.load(
                    file
                )

            text = json.dumps(
                data,
                ensure_ascii=False,
                indent=2,
            )

            return ExtractedDocument(
                text=text,
                parts=[
                    ExtractedPart(
                        part_type="json",
                        name="JSON",
                        text=text,
                    )
                ],
                parser_name="json",
                parser_version="stdlib",
                status="ok",
            )

        except Exception as exc:
            last_error = exc

    return ExtractedDocument(
        status="error",
        error=str(last_error),
        parser_name="json",
    )


# ============================================================
# DWG
# ============================================================

def extract_dwg(
    path: Path,
) -> ExtractedDocument:

    """
    На первом этапе DWG НЕ разбираем как CAD-графику.

    Мы не будем придумывать содержимое чертежа.

    Индексируем только безопасные данные:
    имя файла, путь, расширение.

    Полноценный DWG-парсер можно подключить позже.
    """

    text = (
        f"DWG файл: {path.name}\n"
        f"Путь: {path}\n"
        f"Формат: AutoCAD DWG\n"
        f"Содержимое геометрии на данном этапе "
        f"не извлекалось."
    )

    return ExtractedDocument(
        text=text,
        parts=[
            ExtractedPart(
                part_type="dwg_metadata",
                name=path.name,
                text=text,
            )
        ],
        parser_name="dwg-metadata",
        parser_version="1.0",
        status="metadata_only",
        metadata={
            "geometry_parsed": False,
        },
    )


# ============================================================
# DOC
# ============================================================

def _doc_pieces_text(word: bytes, table: bytes) -> str | None:
    """
    Текст из Word 97-2003 (.doc) по таблице фрагментов (Piece Table, CLX).
    Возвращает None, если структура не распознана.
    """

    import struct

    if len(word) < 0x1AA or struct.unpack_from("<H", word, 0)[0] != 0xA5EC:
        return None

    ccp_text = struct.unpack_from("<i", word, 0x4C)[0]
    fc_clx, lcb_clx = struct.unpack_from("<II", word, 0x1A2)

    if ccp_text <= 0 or lcb_clx == 0 or fc_clx + lcb_clx > len(table):
        return None

    clx = table[fc_clx:fc_clx + lcb_clx]

    pos = 0

    # Пропускаем блоки Prc (0x01), ищем Pcdt (0x02).
    while pos < len(clx) and clx[pos] == 0x01:
        cb = struct.unpack_from("<H", clx, pos + 1)[0]
        pos += 3 + cb

    if pos >= len(clx) or clx[pos] != 0x02:
        return None

    lcb = struct.unpack_from("<I", clx, pos + 1)[0]
    plc = clx[pos + 5:pos + 5 + lcb]

    count = (len(plc) - 4) // 12

    if count <= 0:
        return None

    cps = struct.unpack_from(f"<{count + 1}I", plc, 0)

    chunks: list[str] = []

    for i in range(count):

        cp_start, cp_end = cps[i], cps[i + 1]

        if cp_start >= ccp_text:
            break

        cp_end = min(cp_end, ccp_text)

        fc = struct.unpack_from("<I", plc, (count + 1) * 4 + i * 8 + 2)[0]

        length = cp_end - cp_start

        if fc & 0x40000000:
            offset = (fc & 0x3FFFFFFF) // 2
            raw = word[offset:offset + length]
            chunks.append(raw.decode("cp1252", "replace"))
        else:
            offset = fc & 0x3FFFFFFF
            raw = word[offset:offset + length * 2]
            chunks.append(raw.decode("utf-16-le", "replace"))

    return "".join(chunks)


def _doc_clean(text: str) -> str:

    import re

    # Поля Word: убираем коды (\x13 инструкция \x14 результат \x15).
    text = re.sub(r"\x13[^\x13\x14\x15]*[\x14\x15]", "", text)

    text = (
        text.replace("\r", "\n")
        .replace("\x07", "\t")
        .replace("\x0b", "\n")
        .replace("\x0c", "\n")
    )

    text = re.sub(r"[\x00-\x08\x0e-\x1f]", "", text)

    return _clean_text(text)


def _doc_strings_fallback(word: bytes) -> str:
    """
    Запасной вариант: достаём читаемые строки из потока WordDocument
    (UTF-16 и 8-bit). Хуже, чем разбор структуры, но лучше, чем ничего.
    """

    import re

    found: list[str] = []

    for m in re.finditer(rb"(?:[\x20-\x7e\r\n\t]\x00){6,}", word):
        found.append(m.group().decode("utf-16-le", "ignore"))

    for m in re.finditer(rb"[\x20-\x7e\xc0-\xff\r\n\t]{12,}", word):
        found.append(m.group().decode("cp1251", "ignore"))

    for m in re.finditer(rb"(?:[\x00-\x04][\x04]|[\x20-\x7e]\x00|[\x10-\x4f]\x04){6,}", word):
        found.append(m.group().decode("utf-16-le", "ignore"))

    return "\n".join(found)


def extract_doc(
    path: Path,
) -> ExtractedDocument:
    """
    Старый .DOC (Word 97-2003). Читается без Word и LibreOffice:
    контейнер OLE открывается пакетом olefile, текст собирается из
    таблицы фрагментов. Если формат не распознан — запасной разбор
    строк, а если и он ничего не дал — запись только с метаданными.
    """

    text = ""
    parser_name = "doc-olefile"
    note = ""

    try:

        import olefile

        ole = olefile.OleFileIO(str(path))

        try:

            word = ole.openstream("WordDocument").read()

            table_name = "0Table"

            if len(word) > 11 and (word[0x0B] & 0x02):
                table_name = "1Table"

            table = b""

            if ole.exists(table_name):
                table = ole.openstream(table_name).read()

            try:
                raw_text = _doc_pieces_text(word, table)
            except Exception:
                raw_text = None

            if raw_text:
                text = _doc_clean(raw_text)
            else:
                text = _doc_clean(_doc_strings_fallback(word))
                parser_name = "doc-strings"
                note = "структура не распознана, текст извлечён приблизительно"

        finally:
            ole.close()

    except ImportError:
        note = "пакет olefile не установлен"

    except Exception as exc:
        note = f"не удалось прочитать файл: {exc}"

    if len(text.strip()) >= 20:

        return ExtractedDocument(
            text=text,
            parts=[
                ExtractedPart(
                    part_type="text",
                    name=path.name,
                    text=text,
                )
            ],
            parser_name=parser_name,
            parser_version="1.0",
            status="ok",
            metadata={"doc_note": note} if note else {},
        )

    meta_text = (
        f"DOC файл: {path.name}\n"
        f"Путь: {path}\n"
        f"Формат: Microsoft Word DOC\n"
        f"Текст не извлечён"
        + (f" ({note})." if note else ".")
    )

    return ExtractedDocument(
        text=meta_text,
        parts=[
            ExtractedPart(
                part_type="doc_metadata",
                name=path.name,
                text=meta_text,
            )
        ],
        parser_name="doc-metadata",
        parser_version="1.0",
        status="metadata_only",
    )


# ============================================================
# Универсальный extractor
# ============================================================

EXTRACTORS = {
    ".pdf": extract_pdf,
    ".doc": extract_doc,
    ".docx": extract_docx,
    ".xls": extract_xls,
    ".xlsx": extract_excel,
    ".xlsm": extract_excel,
    ".xml": extract_xml,
    ".txt": extract_txt,
    ".csv": extract_csv,
    ".json": extract_json,
    ".dwg": extract_dwg,
}


def extract_document(
    path: str | Path,
) -> ExtractedDocument:

    path = Path(path)

    extension = (
        path.suffix.lower()
    )

    extractor = EXTRACTORS.get(
        extension
    )

    if extractor is None:

        return ExtractedDocument(
            status="unsupported",
            error=(
                f"Формат {extension or '(без расширения)'} "
                f"не поддерживается."
            ),
            parser_name="none",
        )

    return extractor(
        path
    )